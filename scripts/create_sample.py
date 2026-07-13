from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Font, PatternFill
from PIL import Image as PillowImage, ImageDraw


root = Path(__file__).resolve().parents[1]
samples = root / "samples"; samples.mkdir(exist_ok=True)
photo = samples / "defect_sample.png"
canvas = PillowImage.new("RGB", (220, 110), "#d9e4f5")
draw = ImageDraw.Draw(canvas); draw.rectangle((45, 22, 175, 88), fill="#707f98"); draw.ellipse((95, 44, 126, 75), fill="#e33f43"); canvas.save(photo)
wb = Workbook(); ws = wb.active; ws.title = "Defects"
ws.append(["Factory Defect Register"]); ws.merge_cells("A1:F1"); ws["A1"].font = Font(bold=True, size=16)
ws.append([]); ws.append(["Defect ID", "Date", "Model", "Quantity", "Cost", "Photo"])
header_fill = PatternFill("solid", fgColor="2457C5")
for cell in ws[3]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = header_fill
ws.append(["D-001", "2026-07-01", "TV-A", 2, 125.50, "screen mark"])
ws.append(["D-002", "2026-07-01", "TV-A", 3, 210.00, "panel scratch"])
img = Image(photo); img.width = 110; img.height = 55; ws.add_image(img, "F5"); ws.row_dimensions[5].height = 45
ws.append(["Subtotal", None, None, 5, 335.50, None])
ws.append(["Mobile line"]); ws.merge_cells("A7:F7")
ws.append(["D-003", "2026-07-02", "M-A", 1, 80.00, "housing dent"])
ws.append(["D-004", "2026-07-02", "M-B", 4, 160.00, "paint defect"])
ws.append(["Subtotal", None, None, 5, 240.00, None])
ws.append(["Grand Total", None, None, 10, 575.50, None])
ws.column_dimensions["A"].width = 18; ws.column_dimensions["B"].width = 16; ws.column_dimensions["C"].width = 14; ws.column_dimensions["F"].width = 24
target = samples / "factory_defects.xlsx"; wb.save(target); print(target)

