from __future__ import annotations

import hashlib
import re
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .classifier import classify_row
from .models import ImageAsset, ParsedRow, ParsedSheet, Reconciliation


def _safe_header(value: Any, column: int, seen: dict[str, int]) -> str:
    base = re.sub(r"\s+", " ", str(value).strip()) if value not in (None, "") else f"column_{column}"
    seen[base] += 1
    return base if seen[base] == 1 else f"{base}_{seen[base]}"


def _find_header_row(ws, scan_rows: int = 40) -> int:
    best_row, best_score = 1, float("-inf")
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, scan_rows)):
        values = [c.value for c in row]
        populated = [v for v in values if v not in (None, "")]
        if not populated:
            continue
        strings = sum(isinstance(v, str) for v in populated)
        unique = len({str(v).strip().lower() for v in populated})
        formulas = sum(isinstance(v, str) and v.startswith("=") for v in populated)
        score = len(populated) * 2 + strings + unique * 0.25 - formulas * 3
        if len(populated) >= 2 and score > best_score:
            best_row, best_score = row[0].row, score
    return best_row


def _merged_rows(ws) -> set[int]:
    rows: set[int] = set()
    for cell_range in ws.merged_cells.ranges:
        rows.update(range(cell_range.min_row, cell_range.max_row + 1))
    return rows


def _extract_images(ws, images_dir: Path) -> list[ImageAsset]:
    assets: list[ImageAsset] = []
    images_dir.mkdir(parents=True, exist_ok=True)
    for image in getattr(ws, "_images", []):
        try:
            raw = image._data()
            anchor = image.anchor._from
            row, column = anchor.row + 1, anchor.col + 1
            extension = (getattr(image, "format", None) or "png").lower()
            if extension == "jpeg":
                extension = "jpg"
            digest = hashlib.sha256(raw).hexdigest()
            target = images_dir / f"{digest}.{extension}"
            if not target.exists():
                target.write_bytes(raw)
            assets.append(ImageAsset(
                sheet=ws.title, source_row=row, source_column=column,
                anchor_cell=f"{get_column_letter(column)}{row}", sha256=digest,
                path=target, extension=extension,
            ))
        except Exception:
            # A corrupt drawing must not destroy otherwise valid tabular ingestion.
            continue
    return assets


def _number(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        cleaned = value.replace(",", "").strip()
        try:
            return float(cleaned)
        except ValueError:
            return None
    return None


def _reconcile(sheet: str, rows: list[ParsedRow], headers: list[str]) -> list[Reconciliation]:
    output: list[Reconciliation] = []
    all_data: list[ParsedRow] = []
    section_data: list[ParsedRow] = []
    for row in rows:
        if row.row_type == "section":
            section_data = []
        elif row.row_type == "data":
            all_data.append(row)
            section_data.append(row)
        elif row.row_type in ("total", "subtotal"):
            scope = "all_preceding_data" if "grand_total_label" in row.reasons else "current_section"
            source = all_data if scope == "all_preceding_data" else section_data
            for header in headers:
                written = _number(row.values.get(header))
                values = [_number(r.values.get(header)) for r in source]
                numeric = [v for v in values if v is not None]
                if written is None or not numeric:
                    continue
                calculated = round(sum(numeric), 8)
                difference = round(written - calculated, 8)
                tolerance = max(0.01, abs(written) * 0.000001)
                output.append(Reconciliation(
                    sheet=sheet, source_row=row.source_row, column=header,
                    written_value=written, calculated_value=calculated,
                    difference=difference, status="pass" if abs(difference) <= tolerance else "fail",
                    scope=scope,
                ))
            if scope == "current_section":
                section_data = []
    return output


def parse_workbook(path: Path, images_dir: Path) -> list[ParsedSheet]:
    keep_vba = path.suffix.lower() == ".xlsm"
    wb = load_workbook(path, read_only=False, data_only=False, keep_vba=keep_vba)
    wb_values = load_workbook(path, read_only=False, data_only=True, keep_vba=keep_vba)
    parsed: list[ParsedSheet] = []
    try:
        for ws in wb.worksheets:
            values_ws = wb_values[ws.title]
            header_row = _find_header_row(ws)
            seen: dict[str, int] = defaultdict(int)
            headers = [_safe_header(ws.cell(header_row, col).value, col, seen) for col in range(1, ws.max_column + 1)]
            merged_rows = _merged_rows(ws)
            images = _extract_images(ws, images_dir)
            images_by_row: dict[int, list[str]] = defaultdict(list)
            for asset in images:
                images_by_row[asset.source_row].append(asset.sha256)
            rows: list[ParsedRow] = []
            for row_number in range(header_row + 1, ws.max_row + 1):
                formula_values = [ws.cell(row_number, c).value for c in range(1, ws.max_column + 1)]
                cached_values = [values_ws.cell(row_number, c).value for c in range(1, ws.max_column + 1)]
                effective = [cached_values[i] if cached_values[i] is not None else formula_values[i] for i in range(len(headers))]
                result = classify_row(effective, formula_values, merged=row_number in merged_rows)
                rows.append(ParsedRow(
                    source_row=row_number, row_type=result.row_type,
                    values=dict(zip(headers, effective)), reasons=result.reasons,
                    image_hashes=images_by_row.get(row_number, []),
                ))
            reconciliations = _reconcile(ws.title, rows, headers)
            parsed.append(ParsedSheet(ws.title, header_row, headers, rows, images, reconciliations))
    finally:
        wb.close()
        wb_values.close()
    return parsed

